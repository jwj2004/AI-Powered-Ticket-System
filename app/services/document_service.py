import re

from sqlalchemy import inspect, text
from sqlalchemy.orm import Session

from app.core.database import engine
from app.core.logger import log
from app.models.doc_space import DocSpace
from app.models.document import Document
from app.models.document_chunk import DocumentChunk
from app.models.doc_version import DocVersion

_FAQ_TITLE_RE = re.compile(r"FAQ|常见问题|新手指南")
_FAQ_Q_RE = re.compile(r"(?m)^#{2,3}[ \t]*Q:")


def split_text(text: str, chunk_size: int = 500, overlap: int = 50) -> list[str]:
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start = end - overlap
    return [c for c in chunks if c.strip()]


def is_faq_document(title: str | None, content: str) -> bool:
    if title and _FAQ_TITLE_RE.search(title):
        return True
    return bool(_FAQ_Q_RE.search(content or ""))


def split_faq(content: str, title: str | None) -> list[str]:
    """按 ## Q: / ### Q: 切成单条问答，块首保留文档标题。"""
    matches = list(_FAQ_Q_RE.finditer(content or ""))
    if not matches:
        return []
    prefix = (title or "").strip()
    chunks = []
    for i, match in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(content)
        body = content[match.start():end].strip()
        if not body:
            continue
        chunks.append(f"{prefix}\n\n{body}" if prefix else body)
    return chunks


def split_document(content: str, title: str | None = None) -> list[tuple[str, str]]:
    """返回 (块文本, chunk_type)。FAQ 按单条问答切，其余仍按 500/50。"""
    if is_faq_document(title, content):
        parts = split_faq(content, title)
        if parts:
            return [(part, "faq") for part in parts]
    return [(part, "normal") for part in split_text(content or "")]


def ensure_chunk_type_column() -> None:
    """旧库没有 chunk_type 时补列，避免查询失败。"""
    cols = {col["name"] for col in inspect(engine).get_columns("document_chunk")}
    if "chunk_type" in cols:
        return
    with engine.begin() as conn:
        conn.execute(text("ALTER TABLE document_chunk ADD COLUMN chunk_type VARCHAR(16)"))


def parse_pdf(file_path: str) -> str:
    from PyPDF2 import PdfReader
    reader = PdfReader(file_path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def parse_docx(file_path: str) -> str:
    import docx
    doc = docx.Document(file_path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def parse_md(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()


def parse_txt(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()


PARSERS = {
    "pdf": parse_pdf,
    "docx": parse_docx,
    "md": parse_md,
    "txt": parse_txt,
}


def parse_file(file_path: str, content_type: str) -> str:
    parser = PARSERS.get(content_type)
    if not parser:
        raise ValueError(f"不支持的文件类型: {content_type}")
    return parser(file_path)


def create_document(
    db: Session,
    space_id: int,
    title: str,
    content: str,
    content_type: str,
    owner_id: int,
) -> Document:
    doc = Document(
        space_id=space_id,
        title=title,
        content=content,
        content_type=content_type,
        version=1,
        owner_id=owner_id,
    )
    db.add(doc)
    db.flush()

    version = DocVersion(
        document_id=doc.id,
        version=1,
        content=content,
    )
    db.add(version)

    ensure_chunk_type_column()
    chunks = split_document(content, title)
    chunk_ids = []
    for i, (chunk_text, chunk_type) in enumerate(chunks):
        chunk = DocumentChunk(
            document_id=doc.id,
            chunk_index=i,
            content=chunk_text,
            chunk_type=chunk_type,
        )
        db.add(chunk)
        db.flush()
        chunk_ids.append(chunk.id)

    db.commit()
    db.refresh(doc)

    # 向量化入 FAISS 索引
    if chunks:
        from app.services.vector_retriever import VectorRetriever
        VectorRetriever.add_document_chunks([text for text, _ in chunks], chunk_ids)

    log.info(f"文档创建: {title} (space={space_id}, chunks={len(chunks)})")
    return doc


def delete_document(db: Session, doc_id: int):
    chunks = db.query(DocumentChunk).filter(DocumentChunk.document_id == doc_id).all()
    chunk_ids = [c.id for c in chunks]

    db.query(DocumentChunk).filter(DocumentChunk.document_id == doc_id).delete()
    db.query(DocVersion).filter(DocVersion.document_id == doc_id).delete()
    db.query(Document).filter(Document.id == doc_id).delete()
    db.commit()

    if chunk_ids:
        from app.services.vector_retriever import VectorRetriever
        VectorRetriever.remove_document_chunks(chunk_ids)

    log.info(f"文档删除: id={doc_id}")


def update_document(db: Session, doc_id: int, content: str, title: str | None = None):
    doc = db.query(Document).filter(Document.id == doc_id).first()
    if not doc:
        raise ValueError("文档不存在")

    old_chunks = db.query(DocumentChunk).filter(DocumentChunk.document_id == doc_id).all()
    old_chunk_ids = [c.id for c in old_chunks]

    doc.version += 1
    doc.content = content
    if title is not None:
        doc.title = title

    version = DocVersion(
        document_id=doc.id,
        version=doc.version,
        content=content,
    )
    db.add(version)

    db.query(DocumentChunk).filter(DocumentChunk.document_id == doc_id).delete()
    ensure_chunk_type_column()
    chunks = split_document(content, doc.title)
    chunk_ids = []
    for i, (chunk_text, chunk_type) in enumerate(chunks):
        chunk = DocumentChunk(
            document_id=doc_id,
            chunk_index=i,
            content=chunk_text,
            chunk_type=chunk_type,
        )
        db.add(chunk)
        db.flush()
        chunk_ids.append(chunk.id)

    db.commit()
    db.refresh(doc)

    # 重建该文档的向量索引
    if old_chunk_ids:
        from app.services.vector_retriever import VectorRetriever
        VectorRetriever.remove_document_chunks(old_chunk_ids)
    if chunks:
        from app.services.vector_retriever import VectorRetriever
        VectorRetriever.add_document_chunks([text for text, _ in chunks], chunk_ids)

    log.info(f"文档更新: {doc.title} v{doc.version} (chunks={len(chunks)})")
    return doc


def rollback_document(db: Session, doc_id: int, target_version: int):
    doc = db.query(Document).filter(Document.id == doc_id).first()
    if not doc:
        raise ValueError("文档不存在")

    target = (
        db.query(DocVersion)
        .filter(DocVersion.document_id == doc_id, DocVersion.version == target_version)
        .first()
    )
    if not target:
        raise ValueError(f"版本 {target_version} 不存在")

    old_chunks = db.query(DocumentChunk).filter(DocumentChunk.document_id == doc_id).all()
    old_chunk_ids = [c.id for c in old_chunks]

    doc.version += 1
    doc.content = target.content

    new_version = DocVersion(
        document_id=doc.id,
        version=doc.version,
        content=target.content,
    )
    db.add(new_version)

    db.query(DocumentChunk).filter(DocumentChunk.document_id == doc_id).delete()
    ensure_chunk_type_column()
    chunks = split_document(target.content, doc.title)
    chunk_ids = []
    for i, (chunk_text, chunk_type) in enumerate(chunks):
        chunk = DocumentChunk(
            document_id=doc_id,
            chunk_index=i,
            content=chunk_text,
            chunk_type=chunk_type,
        )
        db.add(chunk)
        db.flush()
        chunk_ids.append(chunk.id)

    db.commit()
    db.refresh(doc)

    if old_chunk_ids:
        from app.services.vector_retriever import VectorRetriever
        VectorRetriever.remove_document_chunks(old_chunk_ids)
    if chunks:
        from app.services.vector_retriever import VectorRetriever
        VectorRetriever.add_document_chunks([text for text, _ in chunks], chunk_ids)

    log.info(f"文档回滚: {doc.title} v{doc.version} (回滚至 v{target_version})")
    return doc
